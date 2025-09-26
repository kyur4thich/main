import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

# Load original file
source_file = "/Users/nic/Downloads/Second Fuyo Q&Asheet.xlsx"
output_file = "/Users/nic/Downloads/Combined_Fuyo_QAs.xlsx"

headers = ["NO", "Related Document", "page", "Date", "Question(Japanese)", "Question(English)", "Answer"]

# Read and combine data
wb_src = openpyxl.load_workbook(source_file, data_only=True)
all_rows = []

for sheet_name in wb_src.sheetnames:
    ws = wb_src[sheet_name]
    for row in ws.iter_rows(min_row=4, min_col=2, max_col=8, values_only=True):
        if any(cell is not None for cell in row):
            all_rows.append(row)

df = pd.DataFrame(all_rows, columns=headers)

# Save to Excel (raw)
df.to_excel(output_file, index=False, sheet_name="Combined")

# Apply formatting
wb = openpyxl.load_workbook(output_file)
ws = wb["Combined"]

# Set column widths
ws.column_dimensions["B"].width = 20  # Related Document
ws.column_dimensions["E"].width = 40  # Question (JP)
ws.column_dimensions["F"].width = 40  # Question (EN)
ws.column_dimensions["G"].width = 90  # Answer

# Format header row A1 to G1
header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
arial_font = Font(name="Arial", bold=True)

for col in range(1, 8):  # A to G = 1 to 7
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = arial_font
    cell.alignment = Alignment(wrap_text=True)

# Wrap text + Arial font in all cells
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=7):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True)
        cell.font = Font(name="Arial")

# Save final version
wb.save(output_file)

print(f"Formatted Excel saved to: {output_file}")
