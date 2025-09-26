from boxsdk import JWTAuth, Client
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink
import os
import time
import traceback
import sys
from datetime import datetime

# === Box Auth ===
auth = JWTAuth.from_settings_file('/Users/nic/code/box-scripts/config/976402037_13exz3x7_config.json')
client = Client(auth)

# === Folder Map ===
folder_map = {
    "01": {"id": "316984976711", "name": "DD Working Folder"},
    "02": {"id": "283328524100", "name": "Wilhelmina Project Data Room - Confidential"},
    "03": {"id": "322833268600", "name": "Wilhelmina TG2 10TPH - Kuantan #1 - Shared Project Files"},
}

# === Prompt selection ===
print("📁 Available folders:")
for key, val in folder_map.items():
    print(f"  {key}: {val['name']}")
selected = input("🔹 Select folder number to analyze (01, 02, 03): ").zfill(2)
if selected not in folder_map:
    print("❌ Invalid selection.")
    sys.exit(1)

selected_folder = folder_map[selected]
records = []

# === Recursive file walker ===
FIELDS = ['type', 'id', 'name', 'size', 'created_at', 'modified_at']

def walk(folder_id, path=""):
    try:
        offset = 0
        while True:
            items = client.folder(folder_id).get_items(limit=1000, offset=offset, fields=FIELDS)
            page_count = 0
            for item in items:
                page_count += 1
                if item.type == "folder":
                    walk(item.id, os.path.join(path, item.name))
                elif item.type == "file":
                    try:
                        size_kb = round((item.size or 0) / 1024, 2)

                        # Split the path and extract first subfolder beneath root
                        parts = path.split(os.sep)
                        folder_name = parts[1] if len(parts) > 1 else ""  # Subfolder1 if exists

                        records.append({
                            "Date Created": getattr(item, "created_at", None),
                            "Date Modified": getattr(item, "modified_at", None),
                            "File Name": item.name,
                            "File Location": os.path.join(path, item.name),
                            "Folder": folder_name,   # first subfolder under the root
                            "File Size (KB)": size_kb,
                            "View File": f"https://app.box.com/file/{item.id}",
                        })
                        if len(records) % 100 == 0:
                            print(f"✅ Processed {len(records)} files...")
                    except Exception as e:
                        print(f"⚠️ Error recording file {item.name} (ID: {item.id}): {e}")
                        traceback.print_exc()
                        continue
            if page_count < 1000:
                break
            offset += page_count
            time.sleep(0.2)
    except KeyboardInterrupt:
        print("\n⛔ Interrupted. Saving what we have...")
        raise
    except Exception as e:
        print(f"❌ Folder error at {path} (ID: {folder_id}): {e}")
        traceback.print_exc()

# === Start walk ===
print(f"\n🔎 Walking: {selected_folder['name']}\n")
try:
    walk(selected_folder["id"], selected_folder["name"])
except KeyboardInterrupt:
    pass  # still export below

# === To Excel ===
df = pd.DataFrame(records)
if not df.empty and "Date Modified" in df.columns:
    df.sort_values(by="Date Modified", ascending=False, inplace=True)

# Format today's date as YYYYMMDD   
today_str = datetime.today().strftime("%Y%m%d")

output_path = f"/Users/nic/Desktop/Data_Room_Inventory_{selected_folder['name'].replace(' ', '_')}_{today_str}.xlsx"


wb = Workbook()
ws = wb.active
ws.title = "Inventory"

# Header
cols = list(df.columns) if not df.empty else [
    "Date Created","Date Modified","File Name",
    "File Location","Folder","File Size (KB)","View File"
]
for col_idx, col_name in enumerate(cols, start=1):
    ws.cell(row=1, column=col_idx, value=col_name)

# Rows
for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
    for col_idx, col_name in enumerate(cols, start=1):
        value = row.get(col_name, "")
        cell = ws.cell(row=row_idx, column=col_idx)
        if col_name == "View File" and isinstance(value, str) and value.startswith("http"):
            cell.value = "🔗 View File"
            cell.hyperlink = Hyperlink(ref=cell.coordinate, target=value)
            cell.font = Font(color="0563C1", underline="single")
            cell.alignment = Alignment(horizontal="left")
        else:
            cell.value = value

# Autofit
for col in ws.columns:
    max_len = 0
    col_letter = col[0].column_letter
    for cell in col:
        if cell.value is not None:
            try:
                max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
    ws.column_dimensions[col_letter].width = min(max_len + 5, 80)

wb.save(output_path)
print(f"\n✅ Export complete: {output_path}")
