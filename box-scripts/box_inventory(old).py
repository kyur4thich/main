from boxsdk import JWTAuth, Client
import pandas as pd
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink
from datetime import datetime

# === Authenticate using Box JWT config ===
auth = JWTAuth.from_settings_file('/Users/nic/code/box-scripts/config/976402037_13exz3x7_config.json')
client = Client(auth)

me = client.user().get()
print(f"🔐 Authenticated as: {me.name} ({me.id})")
print(f"📧 App User login: {me.login}")

# === Force-accept pending collaboration invites via raw API ===
print("📬 Force-checking for pending invites via direct API...")
url = 'https://api.box.com/2.0/collaborations?status=pending'
headers = {'Content-Type': 'application/json'}
response = client._session.get(url, headers=headers)
pending = response.json()

accepted_any = False
for collab in pending['entries']:
    collab_id = collab['id']
    folder_name = collab['item']['name']
    print(f" - Accepting: {folder_name} (Collaboration ID: {collab_id})")
    client._session.put(
        f'https://api.box.com/2.0/collaborations/{collab_id}',
        headers=headers,
        data=json.dumps({'status': 'accepted'})
    )
    accepted_any = True

if accepted_any:
    print("✅ Pending invites accepted.")
else:
    print("✅ No invites needed accepting.")

# === Folder Map ===
folder_map = {
    "01": {"id": "316984650859", "name": "01. EME Uploads"},
    "02": {"id": "316987304563", "name": "02. WEM Review"},
    "03": {"id": "316982193479", "name": "03. Afry Review"},
    "04": {"id": "316988094360", "name": "04. Frozen Files"},
    "05": {"id": "316985039982", "name": "05. Investor Engineering DD"},
    "10": {"id": "319726840600", "name": "Wilhelmina Project Data Room - Confidential / 10. Engineering"}
}

print("\n📂 Available folders to scan:")
for key, val in folder_map.items():
    print(f"  {key}: {val['name']}")

choice = input("\n🔹 Enter folder number to scan (e.g. 01, 02, 10): ").strip()
folder_info = folder_map.get(choice)
if not folder_info:
    print("❌ Invalid choice. Exiting.")
    exit()

folder_id = folder_info['id']
folder_label = folder_info['name']
folder_code = choice  # Use this directly in the filename

print(f"\n📂 Scanning from: {folder_label} (ID: {folder_id})")

# === Traverse folders and collect file data ===
file_records = []

def traverse_folder(folder_id, path):
    items = client.folder(folder_id).get_items(limit=1000)
    for item in items:
        if item.type == 'folder':
            traverse_folder(item.id, f"{path}/{item.name}")
        elif item.type == 'file':
            file = client.file(item.id).get()
            file_records.append({
                "Folder Name": folder_label,
                "Folder Path": path,
                "File Name": file.name,
                "File ID": file.id,
                "File Link": f"https://app.box.com/file/{file.id}",
                "Size (KB)": round(file.size / 1024, 2),
                "Owner": file.created_by['name'] if file.created_by else None,
                "Uploaded At": file.created_at,
                "Modified At": file.modified_at,
                "Last Downloaded At": "N/A"
            })

print("\n🔍 Scanning Box folders recursively...")
traverse_folder(folder_id, folder_label)

# === Export to Excel ===
df = pd.DataFrame(file_records)


today = datetime.today().strftime('%Y%m%d')
output_path = f"/Users/nic/Desktop/boxVDR_Inventory_{folder_code}_{today}.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "BoxVDR Inventory"

# Header
for col_idx, col_name in enumerate(df.columns, start=1):
    ws.cell(row=1, column=col_idx, value=col_name)

# Rows
for row_idx, (_, row) in enumerate(df.iterrows(), start=2):
    for col_idx, col_name in enumerate(df.columns, start=1):
        value = row[col_name]
        cell = ws.cell(row=row_idx, column=col_idx)
        if col_name == "File Link" and isinstance(value, str) and value.startswith("http"):
            cell.value = "View File"
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=value, target=value)
            cell.font = Font(color="0563C1", underline="single")
            cell.alignment = Alignment(horizontal="left")
        else:
            cell.value = value

# Autofit column width
for col in ws.columns:
    max_len = 0
    col_letter = col[0].column_letter
    for cell in col:
        if cell.value:
            try:
                max_len = max(max_len, len(str(cell.value)))
            except:
                pass
    ws.column_dimensions[col_letter].width = max_len + 5

wb.save(output_path)
print(f"\n✅ Export complete: {output_path}")
